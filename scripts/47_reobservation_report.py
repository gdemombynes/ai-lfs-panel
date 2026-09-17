# ruff: noqa: E501
"""Word report of the re-observation diagnostics (output of 46_reobservation.py)."""

from __future__ import annotations

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from lfspanel.config import OUTPUT, ROOT

NAMES = {
    "BRA": "Brazil",
    "MEX": "Mexico",
    "ZAF": "South Africa",
    "ARG": "Argentina",
    "URY": "Uruguay",
    "BOL": "Bolivia",
    "PER": "Peru",
}
DESIGN = {
    "BRA": "five consecutive quarters; scheduled = interview 1 to 4 (V1016)",
    "MEX": "five consecutive quarters; scheduled = interview 1 to 4 (n_ent)",
    "ZAF": "four consecutive quarters; interview = position in the run of observed quarters; scheduled = position 1 to 3",
    "ARG": "two quarters in, two out, two in; interview = position in the run; scheduled = position 1",
    "URY": "six consecutive monthly interviews (ronda); base = the person's record in the last month of the quarter with ronda 1 to 5, so one interview at least falls in the next quarter; 2021-2022 quarters without ronda excluded",
    "BOL": "four consecutive quarters in numbered rotation groups; interview = position in the run; scheduled = rotation groups other than 0 whose last quarter is later",
    "PER": "two quarters in, two out, two in; interview = position in the run; scheduled = position 1",
}
LABEL = {
    "const": "Constant",
    "male": "Male",
    "interview_2": "Interview 2 (ref. 1)",
    "interview_3": "Interview 3",
    "interview_4": "Interview 4",
    "interview_5": "Interview 5",
    "interview_6": "Interview 6",
    "agegrp_15-21": "Age 15-21 (ref. 30-39)",
    "agegrp_22-25": "Age 22-25",
    "agegrp_26-29": "Age 26-29",
    "agegrp_40-49": "Age 40-49",
    "agegrp_50-64": "Age 50-64",
    "agegrp_65+": "Age 65+",
    "educ_1": "Education: none (ref. primary)",
    "educ_3": "Education: secondary",
    "educ_4": "Education: tertiary",
    "educ_missing": "Education missing",
    "urb_0": "Rural (ref. urban)",
    "urb_missing": "Urban/rural missing",
    "labour_unemployed": "Unemployed (ref. employed, ISCO 5)",
    "labour_inactive": "Inactive",
    "labour_status missing": "Labour status missing",
    "quarter_Q2": "Quarter 2 (ref. Q1)",
    "quarter_Q3": "Quarter 3",
    "quarter_Q4": "Quarter 4",
}


def label(t: str) -> str:
    if t in LABEL:
        return LABEL[t]
    if t.startswith("labour_employed occ "):
        return f"Employed, ISCO major group {t.rsplit(' ', 1)[1]}"
    if t.startswith("year_"):
        return f"Year {t[5:]}"
    return t


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def main() -> None:
    tables = OUTPUT / "tables"
    summ = pd.read_csv(tables / "reobservation_summary.csv").set_index("countrycode")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Inches(0.9))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h = doc.styles["Heading 1"]
    h.font.name = "Calibri"
    h.font.size = Pt(13)
    h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    p = doc.add_paragraph()
    r = p.add_run("Re-observation diagnostics: who is found in the next quarter")
    r.bold = True
    r.font.size = Pt(15)
    doc.add_paragraph(
        "One logit per country, pooling every pair of adjacent quarters in the panel. Sample: persons aged 15 and over in quarter t who are scheduled by the rotation design to be interviewed in t+1. Outcome: the same person id appears in t+1 with the same sex and an age within one year. Unweighted; coefficients are log-odds with standard errors, z and p; odds ratios in the last column. Reference categories: interview 1, age 30-39, primary education, urban, employed in ISCO major group 5 (service and sales), first quarter, first year in the sample. Categories with fewer than 100 persons or with no variation in the outcome are folded into the reference and listed under each table. Samples above 1.2 million person-quarters are randomly subsampled to that size before fitting. Two breaks show up as year effects rather than attrition: South Africa's ids do not link at all between 2025Q4 and 2026Q1 (zero matches, consistent with a new master sample; the 2026 year effect is that pair), and Mexico's link rate falls from about 71 percent to 50 percent for the pairs starting in 2025Q4 and 2026Q1. Those pairs should be excluded from transition estimates until the cause is confirmed with the statistical offices."
    )
    for cc in [c for c in NAMES if c in summ.index]:
        s = summ.loc[cc]
        doc.add_heading(f"{NAMES[cc]}", level=1)
        doc.add_paragraph(
            f"Design: {DESIGN[cc]}. Scheduled person-quarters: {int(s.persons_scheduled):,} over {int(s.quarter_pairs)} quarter pairs; rows fitted: {int(s.rows_fitted):,}. Share found in t+1: {s.share_found:.3f}. McFadden pseudo-R2: {s.pseudo_r2:.3f}. Share found by interview number: {s.found_by_interview}."
            + (
                f" Folded into the reference: {s['folded']}."
                if isinstance(s.get("folded"), str) and s.get("folded")
                else ""
            )
        )
        res = pd.read_csv(tables / f"reobservation_{cc.lower()}.csv")
        t = doc.add_table(rows=1, cols=6)
        t.style = "Table Grid"
        for i, hd in enumerate(["Term", "Coef.", "SE", "z", "p", "Odds ratio"]):
            c = t.rows[0].cells[i]
            shade(c, "E7EEF5")
            rr = c.paragraphs[0].add_run(hd)
            rr.bold = True
            rr.font.size = Pt(8.5)
            if i:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for _, row in res.iterrows():
            cells = t.add_row().cells
            vals = [
                label(row.term),
                f"{row.coef:+.4f}",
                f"{row.se:.4f}",
                f"{row.z:+.2f}",
                f"{row.p:.3f}",
                f"{row.odds_ratio:.3f}",
            ]
            for i, v in enumerate(vals):
                rr = cells[i].paragraphs[0].add_run(v)
                rr.font.size = Pt(8.5)
                if i:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in t.rows:
            row.cells[0].width = Inches(2.6)
            for c in row.cells[1:]:
                c.width = Inches(0.8)
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.space_after = Pt(0)
        doc.add_paragraph()
    out = ROOT / "docs" / "design" / "reobservation-diagnostics.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
